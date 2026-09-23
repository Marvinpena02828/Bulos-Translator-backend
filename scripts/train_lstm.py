"""
LSTM Seq2Seq Translation Model Trainer
=======================================
Trains character-level encoder-decoder LSTM models for all 6 translation
direction pairs:  bul↔en  bul↔tl  en↔tl

Dataset: data/dataset.xlsx
  - COMMON WORDS sheet  : word-level triplets (BULOS, FILIPINO, ENGLISH)
  - WORDS sheet         : additional word-level triplets
  - SENTENCES sheet     : sentence-level triplets (BULOS, TAGALOG, ENGLISH)
                          → 70% training, 30% validation/test split
  - LETTERS sheet       : skipped (alphabet examples, already in alphabet.json)

Each trained model is saved to models/lstm/<src>_to_<tgt>/  with:
  - model.keras          encoder+decoder weights
  - tokenizer.json       character→index vocabulary

Usage:
    python scripts/train_lstm.py [--epochs 100] [--batch 64] [--latent 256]
"""

import argparse
import json
import os
import sys
from pathlib import Path

# ── path setup so we can import project modules if needed ──────────────────
PROJECT_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import numpy as np
import openpyxl
from sklearn.model_selection import train_test_split

os.environ.setdefault("TF_CPP_MIN_LOG_LEVEL", "2")  # suppress TF info logs
import tensorflow as tf
from tensorflow import keras
from tensorflow.keras import layers


# ══════════════════════════════════════════════════════════════════════════════
# 1. DATA LOADING
# ══════════════════════════════════════════════════════════════════════════════

def load_dataset(xlsx_path: str) -> dict:
    """
    Load all sheets from dataset.xlsx and return a dict keyed by language pair.

    Returns:
        {
          "bul_en": [(bul, en), ...],
          "bul_tl": [(bul, tl), ...],
          "en_tl":  [(en, tl), ...],
        }
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    pairs: dict[str, list] = {"bul_en": [], "bul_tl": [], "en_tl": []}

    def clean(v):
        return str(v).strip() if v is not None else ""

    # ── COMMON WORDS: columns BULOS, FILIPINO, ENGLISH ──────────────────────
    ws_words = wb["COMMON WORDS"]
    for row in ws_words.iter_rows(min_row=2, values_only=True):
        bul, tl, en = clean(row[0]), clean(row[1]), clean(row[2])
        if bul and tl and en:
            pairs["bul_tl"].append((bul, tl))
            pairs["bul_en"].append((bul, en))
            pairs["en_tl"].append((en, tl))

    # ── WORDS: columns BULOS, FILIPINO, ENGLISH ──────────────────────────────
    ws_words2 = wb["WORDS "]
    for row in ws_words2.iter_rows(min_row=2, values_only=True):
        bul, tl, en = clean(row[0]), clean(row[1]), clean(row[2])
        if bul and tl and en:
            pairs["bul_tl"].append((bul, tl))
            pairs["bul_en"].append((bul, en))
            pairs["en_tl"].append((en, tl))

    # ── SENTENCES: columns BULOS, TAGALOG, ENGLISH ───────────────────────────
    # 70% training, 30% validation — stored separately
    sentences_raw = []
    ws_sent = wb["SENTENCES"]
    for row in ws_sent.iter_rows(min_row=2, values_only=True):
        bul, tl, en = clean(row[0]), clean(row[1]), clean(row[2])
        if bul and tl and en:
            sentences_raw.append((bul, tl, en))

    sent_train, sent_val = train_test_split(
        sentences_raw, test_size=0.30, random_state=42
    )

    print(f"\nSentences — train: {len(sent_train)}, validation: {len(sent_val)}")

    for bul, tl, en in sent_train:
        pairs["bul_tl"].append((bul, tl))
        pairs["bul_en"].append((bul, en))
        pairs["en_tl"].append((en, tl))

    # Save validation split for evaluation
    val_path = PROJECT_ROOT / "data" / "sentences_validation.json"
    val_path.parent.mkdir(parents=True, exist_ok=True)
    with open(val_path, "w", encoding="utf-8") as f:
        json.dump(
            [{"bulos": b, "tagalog": t, "english": e} for b, t, e in sent_val],
            f, ensure_ascii=False, indent=2
        )
    print(f"Validation split saved → {val_path}")

    for k, v in pairs.items():
        print(f"  {k}: {len(v)} training pairs")

    return pairs


# ══════════════════════════════════════════════════════════════════════════════
# 2. TOKENISER (character-level)
# ══════════════════════════════════════════════════════════════════════════════

START_TOKEN = "\t"   # decoder start-of-sequence
END_TOKEN   = "\n"   # decoder end-of-sequence

def build_tokenizer(texts: list[str]) -> dict:
    """
    Build a character→index vocabulary from a list of strings.
    Index 0 is reserved for padding.
    """
    chars = set()
    for t in texts:
        chars.update(t)
    chars.discard(START_TOKEN)
    chars.discard(END_TOKEN)
    sorted_chars = sorted(chars)
    # 0 = padding, 1 = START_TOKEN, 2 = END_TOKEN, 3+ = regular chars
    char2idx = {START_TOKEN: 1, END_TOKEN: 2}
    for i, c in enumerate(sorted_chars, start=3):
        char2idx[c] = i
    idx2char = {v: k for k, v in char2idx.items()}
    return {"char2idx": char2idx, "idx2char": idx2char}


def encode_sequences(texts: list[str], char2idx: dict, max_len: int,
                     add_start: bool = False, add_end: bool = False) -> np.ndarray:
    """Encode a list of strings to padded integer arrays."""
    seqs = []
    for t in texts:
        seq = []
        if add_start:
            seq.append(char2idx.get(START_TOKEN, 1))
        for c in t:
            seq.append(char2idx.get(c, 0))
        if add_end:
            seq.append(char2idx.get(END_TOKEN, 2))
        # truncate then pad
        seq = seq[:max_len]
        seq += [0] * (max_len - len(seq))
        seqs.append(seq)
    return np.array(seqs, dtype=np.int32)


def encode_target_one_hot(texts: list[str], char2idx: dict,
                          max_len: int, vocab_size: int) -> np.ndarray:
    """
    Encode decoder targets as one-hot shifted by one position
    (teacher-forcing: decoder input has START, target has END).
    """
    out = np.zeros((len(texts), max_len, vocab_size), dtype=np.float32)
    for i, t in enumerate(texts):
        seq = [char2idx.get(c, 0) for c in t]
        seq.append(char2idx.get(END_TOKEN, 2))
        seq = seq[:max_len]
        for j, idx in enumerate(seq):
            if idx != 0:
                out[i, j, idx] = 1.0
    return out


# ══════════════════════════════════════════════════════════════════════════════
# 3. MODEL ARCHITECTURE
# ══════════════════════════════════════════════════════════════════════════════

def build_training_model(src_vocab: int, tgt_vocab: int,
                         latent_dim: int) -> keras.Model:
    """
    Standard seq2seq: LSTM encoder + teacher-forcing LSTM decoder.
    """
    # ── Encoder ──────────────────────────────────────────────────────────────
    enc_inputs = layers.Input(shape=(None,), name="enc_input")
    enc_emb    = layers.Embedding(src_vocab, latent_dim, mask_zero=True,
                                  name="enc_emb")(enc_inputs)
    _, state_h, state_c = layers.LSTM(latent_dim, return_state=True,
                                      name="encoder")(enc_emb)
    enc_states = [state_h, state_c]

    # ── Decoder ──────────────────────────────────────────────────────────────
    dec_inputs = layers.Input(shape=(None,), name="dec_input")
    dec_emb    = layers.Embedding(tgt_vocab, latent_dim, mask_zero=True,
                                  name="dec_emb")(dec_inputs)
    dec_lstm   = layers.LSTM(latent_dim, return_sequences=True,
                              return_state=True, name="decoder")
    dec_out, _, _ = dec_lstm(dec_emb, initial_state=enc_states)
    dec_dense  = layers.Dense(tgt_vocab, activation="softmax", name="dec_dense")
    dec_logits = dec_dense(dec_out)

    model = keras.Model([enc_inputs, dec_inputs], dec_logits)
    return model


def build_inference_models(training_model: keras.Model,
                           latent_dim: int) -> tuple:
    """
    Extract encoder and decoder inference sub-models from the trained model.
    """
    # Encoder inference model
    enc_input  = training_model.get_layer("enc_input").input
    enc_emb    = training_model.get_layer("enc_emb")(enc_input)
    _, h, c    = training_model.get_layer("encoder")(enc_emb)
    enc_model  = keras.Model(enc_input, [h, c])

    # Decoder inference model (stateful)
    dec_input  = training_model.get_layer("dec_input").input
    dec_emb    = training_model.get_layer("dec_emb")(dec_input)
    dec_h_in   = layers.Input(shape=(latent_dim,), name="dec_h_in")
    dec_c_in   = layers.Input(shape=(latent_dim,), name="dec_c_in")
    dec_lstm   = training_model.get_layer("decoder")
    dec_out, dec_h, dec_c = dec_lstm(dec_emb,
                                     initial_state=[dec_h_in, dec_c_in])
    dec_dense  = training_model.get_layer("dec_dense")
    dec_logits = dec_dense(dec_out)
    dec_model  = keras.Model(
        [dec_input, dec_h_in, dec_c_in],
        [dec_logits, dec_h, dec_c]
    )
    return enc_model, dec_model


# ══════════════════════════════════════════════════════════════════════════════
# 4. INFERENCE (greedy decode)
# ══════════════════════════════════════════════════════════════════════════════

def decode_sequence(input_seq: np.ndarray, enc_model: keras.Model,
                    dec_model: keras.Model, tgt_char2idx: dict,
                    tgt_idx2char: dict, max_decode_len: int) -> str:
    states = enc_model.predict(input_seq, verbose=0)
    target_seq = np.array([[tgt_char2idx[START_TOKEN]]])
    result = []
    for _ in range(max_decode_len):
        output, h, c = dec_model.predict(
            [target_seq] + states, verbose=0
        )
        token_idx = np.argmax(output[0, -1, :])
        if token_idx == 0:
            break
        char = tgt_idx2char.get(token_idx, "")
        if char == END_TOKEN:
            break
        result.append(char)
        target_seq = np.array([[token_idx]])
        states = [h, c]
    return "".join(result)


# ══════════════════════════════════════════════════════════════════════════════
# 5. TRAIN ONE DIRECTION PAIR
# ══════════════════════════════════════════════════════════════════════════════

def train_pair(src_lang: str, tgt_lang: str, pairs: list[tuple],
               epochs: int, batch_size: int, latent_dim: int,
               output_dir: Path) -> None:
    print(f"\n{'='*60}")
    print(f"  Training: {src_lang} → {tgt_lang}  ({len(pairs)} pairs)")
    print(f"{'='*60}")

    src_texts = [p[0] for p in pairs]
    tgt_texts = [p[1] for p in pairs]

    # Tokenisers
    src_tok = build_tokenizer(src_texts)
    tgt_tok = build_tokenizer(tgt_texts)
    src_c2i = src_tok["char2idx"]
    tgt_c2i = tgt_tok["char2idx"]
    tgt_i2c = tgt_tok["idx2char"]

    src_vocab = len(src_c2i) + 1   # +1 for padding
    tgt_vocab = len(tgt_c2i) + 1

    max_src = max(len(t) for t in src_texts) + 1
    max_tgt = max(len(t) for t in tgt_texts) + 2   # +2 for START+END

    print(f"  Vocab — src: {src_vocab}, tgt: {tgt_vocab}")
    print(f"  Max len — src: {max_src}, tgt: {max_tgt}")

    # Encode inputs
    enc_in  = encode_sequences(src_texts, src_c2i, max_src)
    dec_in  = encode_sequences(tgt_texts, tgt_c2i, max_tgt, add_start=True)
    dec_out = encode_target_one_hot(tgt_texts, tgt_c2i, max_tgt, tgt_vocab)

    # Build + compile model
    model = build_training_model(src_vocab, tgt_vocab, latent_dim)
    model.compile(
        optimizer=keras.optimizers.Adam(learning_rate=1e-3),
        loss="categorical_crossentropy",
        metrics=["accuracy"]
    )
    model.summary(print_fn=lambda x: None)  # suppress verbose summary

    # Callbacks
    callbacks = [
        keras.callbacks.EarlyStopping(
            monitor="val_loss", patience=10, restore_best_weights=True
        ),
        keras.callbacks.ReduceLROnPlateau(
            monitor="val_loss", factor=0.5, patience=5, min_lr=1e-5
        ),
    ]

    # Train
    history = model.fit(
        [enc_in, dec_in], dec_out,
        batch_size=batch_size,
        epochs=epochs,
        validation_split=0.1,
        callbacks=callbacks,
        verbose=1,
    )

    best_val_loss = min(history.history["val_loss"])
    best_val_acc  = max(history.history.get("val_accuracy",
                                            history.history.get("val_acc", [0])))
    print(f"  Best val_loss: {best_val_loss:.4f}  |  Best val_acc: {best_val_acc:.4f}")

    # Save model weights + tokenisers
    output_dir.mkdir(parents=True, exist_ok=True)
    model.save(output_dir / "model.keras")

    tokenizer_data = {
        "src_lang":   src_lang,
        "tgt_lang":   tgt_lang,
        "max_src_len": max_src,
        "max_tgt_len": max_tgt,
        "src_vocab_size": src_vocab,
        "tgt_vocab_size": tgt_vocab,
        "src_char2idx": src_c2i,
        "tgt_char2idx": tgt_c2i,
        "tgt_idx2char": {str(k): v for k, v in tgt_i2c.items()},
    }
    with open(output_dir / "tokenizer.json", "w", encoding="utf-8") as f:
        json.dump(tokenizer_data, f, ensure_ascii=False, indent=2)

    print(f"  Saved → {output_dir}")

    # Quick smoke-test on a few examples
    enc_inf, dec_inf = build_inference_models(model, latent_dim)
    print("\n  Sample predictions:")
    for src, tgt in pairs[:5]:
        enc_seq = encode_sequences([src], src_c2i, max_src)
        pred    = decode_sequence(enc_seq, enc_inf, dec_inf,
                                  tgt_c2i, tgt_i2c, max_tgt)
        match   = "✓" if pred.strip().lower() == tgt.strip().lower() else " "
        print(f"  {match}  [{src}] → [{pred}]  (expected: [{tgt}])")


# ══════════════════════════════════════════════════════════════════════════════
# 6. MAIN
# ══════════════════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(
        description="Train LSTM seq2seq models for Bulos↔Filipino↔English"
    )
    parser.add_argument("--epochs",  type=int, default=100,
                        help="Max training epochs per direction (default: 100)")
    parser.add_argument("--batch",   type=int, default=64,
                        help="Batch size (default: 64)")
    parser.add_argument("--latent",  type=int, default=256,
                        help="LSTM latent dimension (default: 256)")
    parser.add_argument("--dataset", type=str,
                        default=str(PROJECT_ROOT / "data" / "dataset.xlsx"),
                        help="Path to dataset.xlsx")
    parser.add_argument("--out",     type=str,
                        default=str(PROJECT_ROOT / "models" / "lstm"),
                        help="Output directory for trained models")
    parser.add_argument("--pairs",   type=str, default="all",
                        help="Comma-separated pairs to train, e.g. bul_en,bul_tl "
                             "(default: all)")
    args = parser.parse_args()

    print(f"\nTensorFlow version: {tf.__version__}")
    print(f"Dataset:  {args.dataset}")
    print(f"Output:   {args.out}")
    print(f"Epochs:   {args.epochs}  Batch: {args.batch}  Latent: {args.latent}\n")

    # Load all pairs
    all_pairs = load_dataset(args.dataset)

    # Determine which directions to train
    directions = {
        "bul_en": ("bul", "en"),
        "en_bul": ("en", "bul"),
        "bul_tl": ("bul", "tl"),
        "tl_bul": ("tl", "bul"),
        "en_tl":  ("en", "tl"),
        "tl_en":  ("tl", "en"),
    }

    if args.pairs == "all":
        selected = list(directions.keys())
    else:
        selected = [p.strip() for p in args.pairs.split(",")]

    out_root = Path(args.out)

    for key in selected:
        if key not in directions:
            print(f"Unknown pair '{key}', skipping.")
            continue
        src_lang, tgt_lang = directions[key]

        # Build pairs in the requested direction.
        # Canonical key order: bul < en < tl (alphabetical) so lookups
        # always match what load_dataset() stored.
        base_key = "_".join(sorted([src_lang, tgt_lang]))
        # base_key is always bul_en, bul_tl, or en_tl
        forward_pairs = all_pairs.get(base_key, [])

        if src_lang == base_key.split("_")[0]:
            direction_pairs = forward_pairs
        else:
            direction_pairs = [(b, a) for a, b in forward_pairs]

        if not direction_pairs:
            print(f"No data for {key}, skipping.")
            continue

        train_pair(
            src_lang=src_lang,
            tgt_lang=tgt_lang,
            pairs=direction_pairs,
            epochs=args.epochs,
            batch_size=args.batch,
            latent_dim=args.latent,
            output_dir=out_root / key,
        )

    print("\n\nAll done! Models saved to:", out_root)


if __name__ == "__main__":
    main()
